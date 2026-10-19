"""
processor.py
小分表批量处理逻辑（与界面无关）
"""

import os
import re
import traceback

import pandas as pd
import unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side, PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# ==================== 样式常量 ====================

LIGHT_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")


# ==================== 工具函数 ====================

def clean_header(h):
    """清理表头：删去（共*分）、作答/答案、英文逗号"""
    h = str(h)
    h = re.sub(r"（共\d+分）", "", h)
    h = h.replace("作答/答案", "")
    h = h.replace(",", "")
    return h.strip()


def unique_headers(cols):
    """
    统一表头格式：
    - 纯题号列   "一.1" / "二.16"       → 原样保留
    - 已合并列   "一.1_D" / "二.16_AC"  → "1.D" / "16.AC"
    - 答案列     "C"、"C.1"、"AD.1"     → 挂到上一个题号的小题号，如 "4.C"、"10.AD"
    - 其它列（姓名、总、客、主、校、班等）原样保留
    """
    result = []
    last_q_num = None

    for c in cols:
        c = str(c).strip()

        # 1) 已合并列 "一.1_D" → "1.D"
        m = re.fullmatch(r"[一二三四五六七八九十]+\.(\d+)_([A-Z]+)", c)
        if m:
            q_num, ans = m.group(1), m.group(2)
            last_q_num = q_num
            result.append(f"{q_num}.{ans}")
            continue

        # 2) 纯题号列 "一.1" → 原样保留，记录小题号
        m = re.fullmatch(r"[一二三四五六七八九十]+\.(\d+)", c)
        if m:
            last_q_num = m.group(1)
            result.append(c)
            continue

        # 3) 答案列 "C" / "C.1" / "AD.1" → "小题号.答案"
        m = re.fullmatch(r"([A-Z]+)(?:\.\d+)?", c)
        if m and last_q_num:
            result.append(f"{last_q_num}.{m.group(1)}")
            continue

        # 4) 其它列保持原样
        result.append(c)

    return result


def strip_prefix_from_answer_headers(ws, header_row=2):
    """
    把形如 "1.D"、"16.AC" 的答案列表头去掉题号前缀，只留字母。
    纯题号列（中文开头，如 "一.1"）原样保留。
    """
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row, column=col)
        val = cell.value
        if isinstance(val, str):
            m = re.fullmatch(r"\d+\.([A-Z]+)", val)
            if m:
                cell.value = m.group(1)


def to_num(x):
    """文本形式数字转数字"""
    if isinstance(x, str):
        s = x.strip()
        if s == "":
            return x
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return x


def summarize_zero_scores(df):
    """
    汇总：
    - 除 "二.16" 外的题：筛选得分 == 0 的学生
    - "二.16"：筛选得分 < 0 的学生
    - 排除总分为 0 的学生
    """
    if "姓名" not in df.columns:
        return pd.DataFrame()

    if "总" in df.columns:
        df = df[df["总"] != 0]

    q_cols = [c for c in df.columns
              if re.fullmatch(r"[一二三四五六七八九十]+\.\d+", str(c))]

    rows = []
    for q in q_cols:
        if q == "二.16":
            mask = df[q] < 0
            rule = "< 0"
        else:
            mask = df[q] == 0
            rule = "= 0"

        students = df.loc[mask, "姓名"].dropna().astype(str)
        students = [s.strip() for s in students if s.strip()]

        rows.append({
            "题号": q,
            "筛选规则": rule,
            "人数": len(students),
            "学生名单": "、".join(students) if students else "（无）"
        })

    return pd.DataFrame(rows)


def auto_fit_columns(ws, min_width=4, max_width=60, padding=2):
    """
    根据内容自动调整每列列宽，使其恰好容纳最长内容。
    - 中文/全角字符按 2 计算；英文/数字/半角按 1 计算
    - 用列索引 + get_column_letter 取列名，兼容 MergedCell
    """
    def display_width(s):
        w = 0
        for ch in str(s):
            if unicodedata.east_asian_width(ch) in ("W", "F"):
                w += 2
            else:
                w += 1
        return w

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_w = 0
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is None:
                continue
            for line in str(cell.value).split("\n"):
                max_w = max(max_w, display_width(line))

        width = max(min_width, min(max_w + padding, max_width))
        ws.column_dimensions[col_letter].width = width


# ==================== 单文件处理 ====================

def process_one(input_file, output_file, log=print, exam_name="", class_num=""):
    """处理单个 Excel 文件，输出到 output_file"""
    log(f"正在处理：{os.path.basename(input_file)}")

    # 读取
    df = pd.read_excel(input_file, sheet_name="Worksheet", dtype=object)

    # 1. 删去学号、考号、班级
    for col in ["学号", "考号", "班级"]:
        if col in df.columns:
            df = df.drop(columns=[col])

    # 4. 数据区删英文逗号
    df = df.apply(
        lambda col: col.map(lambda x: x.replace(",", "") if isinstance(x, str) else x)
    )

    # 2、3. 清理表头 + 唯一化
    df.columns = unique_headers([clean_header(c) for c in df.columns])

    # 7. 改列名
    df = df.rename(columns={
        "总分": "总",
        "客观题": "客",
        "主观题": "主",
        "学校排名": "校",
        "班级排名": "班"
    })

    # 5. 文本数字转数字
    df = df.apply(lambda col: col.map(to_num))

    # 6. 删去总分为 0 的行
    if "总" in df.columns:
        df = df[df["总"] != 0].reset_index(drop=True)

    # ===== 写入 Excel =====
    wb = Workbook()
    ws = wb.active
    ws.title = "Worksheet"

    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    # 10. 第一行上方插入空行 → 表头变为第 2 行
    ws.insert_rows(1)

    # 表头行此时是第 2 行：把答案列的 "1.D" / "16.AC" 改成 "D" / "AC"
    header_row = 2
    strip_prefix_from_answer_headers(ws, header_row=header_row)

    max_row = ws.max_row
    max_col = ws.max_column

    # 注意：标题先不写，等列宽调完再写，避免撑宽 A 列

    # ===== 标红 =====
    for row in ws.iter_rows(min_row=header_row + 1, max_row=max_row,
                            min_col=1, max_col=max_col):
        for cell in row:
            col_name = ws.cell(row=header_row, column=cell.column).value
            val = cell.value
            if col_name != "二.16":
                if val == 0:
                    cell.fill = LIGHT_RED
            else:
                if isinstance(val, (int, float)) and val < 0:
                    cell.fill = LIGHT_RED

    # ===== 框线（只覆盖表头到数据区，不含标题行） =====
    for row in ws.iter_rows(min_row=header_row, max_row=max_row,
                            min_col=1, max_col=max_col):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # 外侧加粗
    for col in range(1, max_col + 1):
        c = ws.cell(row=header_row, column=col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=MEDIUM, bottom=c.border.bottom)
        c = ws.cell(row=max_row, column=col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=c.border.top, bottom=MEDIUM)

    for row in range(header_row, max_row + 1):
        c = ws.cell(row=row, column=1)
        c.border = Border(left=MEDIUM, right=c.border.right,
                          top=c.border.top, bottom=c.border.bottom)
        c = ws.cell(row=row, column=max_col)
        c.border = Border(left=c.border.left, right=MEDIUM,
                          top=c.border.top, bottom=c.border.bottom)

    # ===== 列宽自适应（此时 A1 还空着，不受标题影响） =====
    auto_fit_columns(ws)

    # ===== 标题：写在 A1，用“跨列居中”，不合并、不改列宽 =====
    if exam_name or class_num:
        title = f"高三年级{class_num}班{exam_name}考试物理成绩单"
        cell = ws.cell(row=1, column=1)
        cell.value = title
        cell.alignment = Alignment(horizontal="centerContinuous",
                                   vertical="center")
        cell.font = Font(size=14, bold=True)

    # ===== 零分 / 负分汇总工作表 =====
    summary_df = summarize_zero_scores(df)
    if not summary_df.empty:
        ws2 = wb.create_sheet("零分负分汇总")

        for r in dataframe_to_rows(summary_df, index=False, header=True):
            ws2.append(r)

        # 表头加粗
        for cell in ws2[1]:
            cell.font = Font(bold=True)

        # 给"二.16"那行加浅红背景
        for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row,
                                 min_col=1, max_col=4):
            if row[0].value == "二.16":
                for cell in row:
                    cell.fill = LIGHT_RED

        # 表格框线
        for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row,
                                 min_col=1, max_col=4):
            for cell in row:
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

        # 汇总表列宽自适应（放宽一点，方便看名单）
        auto_fit_columns(ws2, max_width=200)

    wb.save(output_file)
    log(f"  ✅ 已生成：{os.path.basename(output_file)}")


# ==================== 批量处理 ====================

def batch_process(input_dir, output_dir, log=print, exts=(".xlsx", ".xls"),
                  exam_name="", class_num=""):
    """
    批量处理 input_dir 下的所有 Excel 文件，输出到 output_dir。
    返回 (成功数, 失败数)。
    """
    os.makedirs(output_dir, exist_ok=True)

    files = [f for f in os.listdir(input_dir)
             if f.lower().endswith(exts) and not f.startswith("~$")]

    if not files:
        log(f"❌ 未在 {input_dir} 找到任何 Excel 文件")
        return 0, 0

    ok, fail = 0, 0
    for f in files:
        input_path = os.path.join(input_dir, f)
        name, ext = os.path.splitext(f)
        output_path = os.path.join(output_dir, f"{name}_处理后{ext}")
        try:
            process_one(input_path, output_path, log,
                        exam_name=exam_name, class_num=class_num)
            ok += 1
        except Exception as e:
            fail += 1
            log(f"  ⚠️ 处理 {f} 失败：{e}")
            log(traceback.format_exc(limit=1))

    log(f"\n完成：成功 {ok} 个，失败 {fail} 个。")
    return ok, fail


def process_class_compare(input_file, output_file, log=print):
    """
    处理班级对比表：
    - 原工作表：文本型数字转数字（A 列“班级”除外）；
      百分比字符串转数字并保留百分比格式
    - 新增工作表“处理后”：班级 + 各题平均分 + 客观分 + 主观分 + 总分
      （班级列保留文本）
    - “客观分”“主观分”两列宽度一致（取较窄者）
    """
    log(f"正在处理班级对比表：{os.path.basename(input_file)}")
    wb = load_workbook(input_file)
    ws_src = wb.active
    rows = list(ws_src.iter_rows(values_only=True))

    # ---------- 找到所有表头行 ----------
    header_indices = []
    for i, row in enumerate(rows):
        if row and len(row) >= 1 and str(row[0]).strip() == "班级":
            if any(cell is not None and "平均分" in str(cell) for cell in row):
                header_indices.append(i)
    if not header_indices:
        raise ValueError("未找到表头行（第一列为“班级”且包含“平均分”的列）")

    # ---------- 解析每个块 ----------
    parsed_blocks = []
    for hi in header_indices:
        header = rows[hi]
        data = []
        j = hi + 1
        while j < len(rows):
            r = rows[j]
            if r is None or all(cell is None or str(cell).strip() == "" for cell in r):
                break
            if j in header_indices:
                break
            if str(r[0]).strip() == "主观题":
                j += 1
                continue
            data.append(r)
            j += 1

        keep_cols = []
        for idx, cell in enumerate(header):
            if cell is None:
                continue
            s = str(cell).strip()
            if s == "班级":
                keep_cols.append((idx, "班级"))
            elif "平均分" in s:
                m = re.search(r'([一二三四五六七八九十]+\.\d+)', s)
                if m:
                    keep_cols.append((idx, m.group(1)))
        if not keep_cols:
            continue

        questions = [name for _, name in keep_cols if name != "班级"]
        if not questions:
            continue

        first_q = questions[0]
        if first_q.startswith("一"):
            block_type = "客观"
        elif first_q.startswith("二"):
            block_type = "主观"
        else:
            block_type = "客观"

        data_dict = {}
        for row in data:
            if row is None:
                continue
            d = {}
            cls_val = None
            for idx, name in keep_cols:
                val = row[idx] if idx < len(row) else None
                if name == "班级":
                    cls_val = val
                    d["班级"] = val
                else:
                    d[name] = to_num(val) if val is not None else None
            if cls_val is None:
                continue
            key = str(cls_val).strip()
            if not key:
                continue
            data_dict[key] = d

        parsed_blocks.append({
            'type': block_type,
            'questions': questions,
            'data': data_dict,
        })

    # ---------- 合并数据 ----------
    class_order = []
    for block in parsed_blocks:
        for cls in block['data'].keys():
            if cls not in class_order:
                class_order.append(cls)

    class_data = {}
    for cls in class_order:
        class_data[cls] = {'客观': {}, '主观': {}, '客观分': 0, '主观分': 0}

    for block in parsed_blocks:
        btype = block['type']
        for cls, d in block['data'].items():
            if cls not in class_data:
                class_data[cls] = {'客观': {}, '主观': {}, '客观分': 0, '主观分': 0}
            for q in block['questions']:
                val = d.get(q)
                class_data[cls][btype][q] = val if isinstance(val, (int, float)) else 0
            total = sum(v for v in class_data[cls][btype].values()
                        if isinstance(v, (int, float)))
            class_data[cls][f"{btype}分"] = total

    obj_questions = []
    subj_questions = []
    for block in parsed_blocks:
        if block['type'] == '客观':
            for q in block['questions']:
                if q not in obj_questions:
                    obj_questions.append(q)
        else:
            for q in block['questions']:
                if q not in subj_questions:
                    subj_questions.append(q)

    # ---------- 写“处理后”工作表（含总分，班级保留文本） ----------
    if "处理后" in wb.sheetnames:
        del wb["处理后"]
    ws_out = wb.create_sheet("处理后")

    headers = (["班级"] + obj_questions + subj_questions
               + ["客观分", "主观分", "总分"])
    ws_out.append(headers)

    for cls in class_order:
        cd = class_data[cls]
        row = [cls]
        for q in obj_questions:
            row.append(cd['客观'].get(q, ""))
        for q in subj_questions:
            row.append(cd['主观'].get(q, ""))
        obj = cd['客观分']
        subj = cd['主观分']
        row.append(obj)
        row.append(subj)
        row.append(obj + subj)
        ws_out.append(row)

    # ---------- 样式 ----------
    for cell in ws_out[1]:
        cell.font = Font(bold=True)

    max_row = ws_out.max_row
    max_col = ws_out.max_column
    for row in ws_out.iter_rows(min_row=1, max_row=max_row,
                                min_col=1, max_col=max_col):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    for col in range(1, max_col + 1):
        c = ws_out.cell(row=1, column=col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=MEDIUM, bottom=c.border.bottom)
        c = ws_out.cell(row=max_row, column=col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=c.border.top, bottom=MEDIUM)

    for r in range(1, max_row + 1):
        c = ws_out.cell(row=r, column=1)
        c.border = Border(left=MEDIUM, right=c.border.right,
                          top=c.border.top, bottom=c.border.bottom)
        c = ws_out.cell(row=r, column=max_col)
        c.border = Border(left=c.border.left, right=MEDIUM,
                          top=c.border.top, bottom=c.border.bottom)

    auto_fit_columns(ws_out)

    # ---------- 让“客观分”和“主观分”列宽一致（取较窄者） ----------
    from openpyxl.utils import get_column_letter
    try:
        obj_idx = headers.index("客观分") + 1
        subj_idx = headers.index("主观分") + 1
        obj_w = ws_out.column_dimensions[get_column_letter(obj_idx)].width or 0
        subj_w = ws_out.column_dimensions[get_column_letter(subj_idx)].width or 0
        target = min(obj_w, subj_w)   # 取较窄者，避免把客观分撑宽
        ws_out.column_dimensions[get_column_letter(obj_idx)].width = target
        ws_out.column_dimensions[get_column_letter(subj_idx)].width = target
    except ValueError:
        pass

    # ---------- 原工作表：文本数字 / 百分比转数字 ----------
    for row in ws_src.iter_rows():
        for cell in row:
            if cell.column == 1:      # 跳过“班级”列
                continue
            val = cell.value
            if not isinstance(val, str):
                continue
            s = val.strip()
            if s == "":
                continue

            # 百分比字符串 → 数字 + 百分比格式
            if s.endswith("%"):
                core = s.rstrip("%")
                try:
                    num = float(core)
                except ValueError:
                    continue
                if "." in core:
                    dec = len(core.split(".")[1])
                    fmt = "0." + ("0" * dec) + "%"
                else:
                    fmt = "0%"
                cell.value = num / 100.0
                cell.number_format = fmt
                continue

            # 普通文本型数字
            new_val = to_num(s)
            if isinstance(new_val, (int, float)):
                cell.value = new_val

    wb.save(output_file)
    log(f"  ✅ 已生成：{os.path.basename(output_file)}")


# ==================== 命令行入口（可选） ====================

if __name__ == "__main__":
    import sys

    in_dir = sys.argv[1] if len(sys.argv) > 1 else "resources"
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "production"
    batch_process(in_dir, out_dir)
