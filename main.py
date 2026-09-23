import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

input_file = "recourses/小分表_150938ced1feec4b.xlsx"
output_file = "recourses/小分表_处理后.xlsx"

# 读取
df = pd.read_excel(input_file, sheet_name="Worksheet", dtype=object)

# 1. 删去“学号”“考号”“班级”三列
df = df.drop(columns=["学号", "考号", "班级"])

# 4. 先把所有英文逗号删掉（数据区）
df = df.apply(
    lambda col: col.map(lambda x: x.replace(",", "") if isinstance(x, str) else x)
)

# 2、3. 清理表头：删分值、删“作答/答案”、删逗号
def clean_header(h):
    h = str(h)
    h = re.sub(r"（共\d+分）", "", h)   # 删去“（共*分）”
    h = h.replace("作答/答案", "")      # 删去“作答/答案”
    h = h.replace(",", "")              # 删英文逗号
    return h.strip()

df.columns = [clean_header(c) for c in df.columns]

# 7. 改列名
df = df.rename(columns={
    "总分": "总",
    "客观题": "客",
    "主观题": "主",
    "学校排名": "校",
    "班级排名": "班"
})

# 5. 文本形式数字转数字
def to_num(x):
    if isinstance(x, str):
        s = x.strip()
        if s == "":
            return x
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return x

df = df.apply(lambda col: col.map(to_num))

# 6. 删去总分为 0 的行
df = df[df["总"] != 0].reset_index(drop=True)

# 写入 Excel：表头先写在第 1 行
wb = Workbook()
ws = wb.active
ws.title = "Worksheet"

for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

# 10. 第一行上方加入空行，所以插入一行，表头变为第 2 行
ws.insert_rows(1)

# 样式参数
thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="000000")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

header_row = 2
max_row = ws.max_row
max_col = ws.max_column

# 8. 标红：除“二.16”外，其余列为 0 标红；“二.16”列小于 0 标红
for row in ws.iter_rows(min_row=header_row + 1, max_row=max_row, min_col=1, max_col=max_col):
    for cell in row:
        col_name = ws.cell(row=header_row, column=cell.column).value
        val = cell.value

        if col_name != "二.16":
            if val == 0:
                cell.fill = red_fill
        else:
            if isinstance(val, (int, float)) and val < 0:
                cell.fill = red_fill

# 9. 整个表格加框线，外侧加粗框线
for row in ws.iter_rows(min_row=header_row, max_row=max_row, min_col=1, max_col=max_col):
    for cell in row:
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

# 外侧加粗：上边
for col in range(1, max_col + 1):
    c = ws.cell(row=header_row, column=col)
    c.border = Border(left=c.border.left, right=c.border.right, top=medium, bottom=c.border.bottom)

# 外侧加粗：下边
for col in range(1, max_col + 1):
    c = ws.cell(row=max_row, column=col)
    c.border = Border(left=c.border.left, right=c.border.right, top=c.border.top, bottom=medium)

# 外侧加粗：左边
for row in range(header_row, max_row + 1):
    c = ws.cell(row=row, column=1)
    c.border = Border(left=medium, right=c.border.right, top=c.border.top, bottom=c.border.bottom)

# 外侧加粗：右边
for row in range(header_row, max_row + 1):
    c = ws.cell(row=row, column=max_col)
    c.border = Border(left=c.border.left, right=medium, top=c.border.top, bottom=c.border.bottom)

wb.save(output_file)
print("已生成：", output_file)
